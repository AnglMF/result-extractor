from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment
from openpyxl.styles import PatternFill

import pandas


class ResultsWorkBook:
    def __init__(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.index = 1

    def new_worksheet(self, name):
        self.index += 1
        return self.workbook.create_sheet(title=name, index=self.index)

    def register_placings(self, info, tournaments):
        self.worksheet = self.new_worksheet("Placings")
        df = pandas.DataFrame.from_dict(info)
        df = df.sort_values("avg_placing")
        orden = ["id", "avg_placing", "win_perc", "name"] + tournaments
        df = df.reindex(columns=orden)
        for r in dataframe_to_rows(df, index=True, header=True):
            if len(r) > 1:
                self.worksheet.append(r)
        for cell in self.worksheet['A'] + self.worksheet[1]:
            cell.style = 'Pandas'

    def rotate_headers(self, sheet, df):
        
        for row in sheet.iter_rows(min_row=1, min_col=3, max_col=len(df.columns)+1, max_row=1):
            for cell in row:
                cell.alignment = Alignment(text_rotation=90)  


    def format_h2h_sheet(self, sheet, columns, df):
        green_fill = PatternFill(start_color="b6d7a8", end_color="b6d7a8", fill_type="solid")
        red_fill = PatternFill(start_color="ea9999", end_color="ea9999", fill_type="solid")
        yellow_fill = PatternFill(start_color="fff2cc", end_color="fff2cc", fill_type="solid")
        gray_fill = PatternFill(start_color="989795", end_color="989795", fill_type="solid")
        for row in sheet.iter_rows(min_row=2, min_col=3, max_row=columns, max_col=len(df.columns)+1):
            for cell in row:
                string1 = f'{cell.column_letter}1'
                string2 = f'B{cell.row}'
                if sheet[string1].value == sheet[string2].value:
                    cell.fill = gray_fill
                if cell.value is not None and isinstance(cell.value, str):
                    if "-" in cell.value:
                        s1, s2 = cell.value.split("-")
                        v1 = int(s1)
                        v2 = int(s2)
                        if v1 > v2:
                            cell.fill = green_fill
                        elif v2 > v1:
                            cell.fill = red_fill
                        else:
                            cell.fill = yellow_fill

        self.rotate_headers(sheet, df) 

    def register_h2h(self, info):
        self.worksheet = self.new_worksheet("H2H")
        df = pandas.DataFrame.from_dict(info)
        player_list = []
        for competitor in info:
            player_list.append(competitor["player"])
        orden = ["player"] + player_list
        df = df.reindex(columns=orden)
        columns = 0
        for r in dataframe_to_rows(df):
            if len(r) > 1:
                self.worksheet.append(r)
                columns+= 1

        self.format_h2h_sheet(self.worksheet, columns, df)


    def register_tournament_results(self, tournament, participants):
        df = pandas.DataFrame.from_dict(participants)
        self.worksheet = self.new_worksheet(tournament)
        orden = ["name", "placing", "seed", "performance"]
        df = df.reindex(columns=orden)
        for r in dataframe_to_rows(df, index=True, header=True):
            if len(r) > 1:
                self.worksheet.append(r)
        self.rotate_headers(self.worksheet, df)

    def register_sets(self, info):
        self.worksheet.title = "Sets Head 2 Head"
        df = pandas.DataFrame.from_dict(info)
        orden = ["score1", "winner", "loser", "score2", "tournament", "round", "result"]
        df = df.reindex(columns=orden)
        for r in dataframe_to_rows(df, index=True, header=True):
            if len(r) > 1:
                self.worksheet.append(r)
        for cell in self.worksheet['A'] + self.worksheet[1]:
            cell.style = 'Pandas'

    def save_workbook(self, name):
        self.workbook.save(name)
        self.workbook.close()
        column = self.worksheet.column_dimensions['B']
        column.alignment = Alignment(horizontal='center')

    def register_games_h2h(self, info):
        self.worksheet = self.new_worksheet("H2H_games")
        df = pandas.DataFrame.from_dict(info)
        player_list = []
        for competitor in info:
                player_list.append(competitor["player"])
        orden = ["player"] + player_list
        df = df.reindex(columns=orden)
        columns = 0
        for r in dataframe_to_rows(df, index=True, header=True):
            if len(r) > 1:
                columns += 1
                self.worksheet.append(r)

        self.format_h2h_sheet(self.worksheet, columns, df)
        return

    def create_spreadsheet(self, ranking_object, placings, tournaments, name="output"):
        self.register_sets(ranking_object.sets.get_sets())
        if ranking_object.ranking.unordered == True:
            self.register_h2h(
                ranking_object.ranking.get_h2h_record(
                        list(ranking_object.ranking.competitors.values())
                    )
                )
            self.register_games_h2h(
                ranking_object.ranking.get_games_h2h_record(
                        list(ranking_object.ranking.competitors.values())
                    )
                )
        else:
            self.register_h2h(
                ranking_object.ranking.get_h2h_record(
                        ranking_object.ranking.competitors_sorted
                    )
                )
            self.register_games_h2h(
                ranking_object.ranking.get_games_h2h_record(
                        ranking_object.ranking.competitors_sorted
                    )
                )
        self.register_placings(placings, tournaments)
        for tournament in tournaments:
            data = ranking_object.ranking.get_single_tournament_results(tournament)
            self.register_tournament_results(tournament, data)
        self.save_workbook(f'{name}.xlsx')